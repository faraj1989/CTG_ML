#!/usr/bin/env python3
"""
CTG ML Pipeline v1 (22-9) - Nested-CV / Leakage-Safe Version
==============================================================
Restructured to address supervisor feedback items 1-11 (see
"التعديلات" review notes). Summary of what changed vs. the original
ctg_pipeline.py:

  1. Single Split Only        -> Repeated Stratified K-Fold CV added on the
                                  80% training portion. The 20% test set is
                                  split off first and is NEVER touched again
                                  until the single final evaluation.
  2. No Confidence Intervals  -> Mean/SD/95% CI reported from the repeated
                                  CV fold scores, plus a bootstrap 95% CI on
                                  the final held-out test evaluation.
  3. No Statistical Comparison-> Paired, corrected-resampled t-test (Nadeau &
                                  Bengio, 2003) and Wilcoxon signed-rank test
                                  between models, using the SAME CV folds.
  4. Best-model-by-Accuracy   -> Primary metrics are PRE-REGISTERED as
                                  Balanced Accuracy + Macro-F1 (set at the
                                  top of this file, before any results are
                                  seen). Accuracy/Precision/Recall/AUC are
                                  still reported as secondary metrics.
  5. "XGBoost not highest AUC"-> Model selection uses only the pre-defined
                                  primary metrics; all metrics are reported
                                  side by side so the trade-off is visible.

  6. SMOTE/SMOTE+Tomek        -> Original vs SMOTE vs SMOTE+Tomek are run as
     undocumented                three explicit, clearly-labeled arms of the
                                  same CV comparison for every model.

  7. Feature selection outside-> Feature selection (if enabled) is a step
     CV (leakage)                inside the imblearn Pipeline, fit ONLY on
                                  each training fold.


  8. Hyperparameter tuning    -> RandomizedSearchCV nested INSIDE each outer
     unclear                     CV training fold (inner CV), never touching
                                  the validation fold or the test set.
  9. "100% leakage-free"      -> Docstrings/report use hedged, precise
                                  language ("... to reduce the risk of data
                                  leakage"), not an absolute claim.

 10. Data availability        -> Left as a single placeholder the author must
     contradiction               fill with ONE true statement (see
                                  DATA_AVAILABILITY_STATEMENT below).

                                  
 11. No External Validation   -> Explicitly reported as a stated limitation
                                  in the generated report, not glossed over.

Golden rule enforced throughout: anything that LEARNS from data (scaling
parameters, selected features, resampling, hyperparameters, model choice)
learns ONLY from the current training fold (during CV) or the 80% training
set (for the final refit) - never from the 20% test set.
"""

import os
import sys
import argparse
import warnings

# Suppress ALL warning messages (sklearn, joblib, xgboost, pandas, ...).
warnings.filterwarnings("ignore")
# The joblib worker processes are separate interpreters and do not see the
# filter above; PYTHONWARNINGS is inherited by them when they start.
os.environ["PYTHONWARNINGS"] = "ignore"

import time
import json
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
import joblib
from datetime import datetime
from scipy import stats
from statsmodels.stats.multitest import multipletests
from sklearn.metrics import make_scorer

from sklearn.model_selection import (train_test_split, RepeatedStratifiedKFold,
                                      StratifiedKFold, RandomizedSearchCV)
from sklearn.preprocessing import StandardScaler, label_binarize
from sklearn.feature_selection import SelectKBest, f_classif
from sklearn.base import clone
from sklearn.ensemble import (RandomForestClassifier, ExtraTreesClassifier,
                              AdaBoostClassifier, GradientBoostingClassifier,
                              StackingClassifier)
from sklearn.linear_model import LogisticRegression
from xgboost import XGBClassifier
from sklearn.metrics import (accuracy_score, precision_score, recall_score,
                             f1_score, balanced_accuracy_score, cohen_kappa_score,
                             matthews_corrcoef, confusion_matrix, classification_report,
                             roc_auc_score)

from imblearn.pipeline import Pipeline as ImbPipeline
from imblearn.over_sampling import SMOTE
from imblearn.combine import SMOTETomek

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

CLASS_NAMES = ["Normal", "Suspect", "Pathologic"]

# ----------------------------------------------------------------------
# PRE-REGISTERED ANALYSIS DECISIONS
# Fix these BEFORE looking at any results (item 4 / item 5 in the review).
# ----------------------------------------------------------------------
PRIMARY_METRICS = ["Balanced_Accuracy", "Macro_F1"]   # used for model selection
PRIMARY_COMPOSITE = "Primary_Selection_Score"
def primary_selection_score(y_true, y_pred):
    """Predefined, equal-weight primary score, used for tuning and ranking."""
    return (balanced_accuracy_score(y_true, y_pred) +
            f1_score(y_true, y_pred, average="macro", zero_division=0)) / 2

PRIMARY_SCORER = make_scorer(primary_selection_score)
SECONDARY_METRICS = ["Accuracy", "Precision_weighted", "Recall_weighted", "AUC_ovr"]

# Item 10: one statement per resource, clearly labeled (matches the manuscript's
# Data Availability section) -- dataset is public, code/models are on request.
DATA_AVAILABILITY_STATEMENT = (
    "Data availability: The CTG dataset used in this study is publicly available from "
    "the UCI Machine Learning Repository. Code and models availability: The complete "
    "pipeline code, saved model bundles, metrics, and plots generated for this study, "
    "together with the companion prediction dashboard, are available from the "
    "corresponding author upon reasonable request."
)

# Item 11: stated explicitly, not glossed over.
LIMITATIONS_STATEMENT = (
    "No external validation dataset was available; therefore, the "
    "generalizability of the findings to independent datasets or different "
    "populations remains to be established."
)

# Item 9: hedged, precise wording instead of an absolute "leakage-free" claim.
METHODOLOGY_LEAKAGE_STATEMENT = (
    "All data-dependent preprocessing, feature selection, resampling, and "
    "hyperparameter tuning procedures were performed strictly within the "
    "training folds of the cross-validation to reduce the risk of data "
    "leakage. The 20% held-out test set was not used in any decision "
    "related to feature selection, resampling, hyperparameter tuning, or "
    "model selection; it was used exactly once, for final evaluation."
)

os.makedirs("results", exist_ok=True)
os.makedirs("plots", exist_ok=True)
os.makedirs("models", exist_ok=True)


# ----------------------------------------------------------------------
# 1. DATA LOADING / LOCKED SPLIT
# ----------------------------------------------------------------------
def load_and_split_data(filepath, test_size=0.20, random_state=42):
    """Load CSV, remove duplicates, and perform ONE stratified 80/20 split.

    The 20% test partition returned here is LOCKED: nothing downstream may
    fit anything on it until the single final evaluation step.
    """
    df = pd.read_csv(filepath, sep=";")
    df.columns = df.columns.str.strip()
    print(f"Loaded {len(df)} rows, {len(df.columns)} columns")

    initial = len(df)
    df = df.drop_duplicates()
    removed = initial - len(df)
    if removed:
        print(f"Removed {removed} duplicate rows")
    print(f"After removing duplicates: {len(df)}")

    X = df.drop(columns=["NSP"])
    y = df["NSP"] - 1  # 1,2,3 -> 0,1,2

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=test_size, random_state=random_state, stratify=y
    )
    print(f"Train (80%): {len(X_train)}   Test (20%, LOCKED until final step): {len(X_test)}")

    nsp = df["NSP"].value_counts()
    print(f"Class distribution - Normal: {nsp.get(1,0)}, Suspect: {nsp.get(2,0)}, Pathologic: {nsp.get(3,0)}")
    return X_train.reset_index(drop=True), X_test.reset_index(drop=True), \
           y_train.reset_index(drop=True), y_test.reset_index(drop=True), X.columns.tolist()


# ----------------------------------------------------------------------
# 2. MODELS + HYPERPARAMETER SEARCH SPACES
# ----------------------------------------------------------------------
def get_base_models(random_state=42):
    """Unfitted estimators; all six support nested hyperparameter search.

    Stacking has internal out-of-fold training (cv=3) in addition to the
    outer evaluation and inner search folds. Its search space is intentionally
    compact to limit computational cost; equal search counts do not imply
    equal training times or identical search-space coverage.
    """
    return {
        # n_jobs=1: the hyperparameter search already uses every core; threads
        # nested inside its workers only add contention (results are unchanged).
        "Random Forest": RandomForestClassifier(n_estimators=300, class_weight="balanced",
                                                 random_state=random_state, n_jobs=1),
        "Extra Trees": ExtraTreesClassifier(n_estimators=300, class_weight="balanced",
                                             random_state=random_state, n_jobs=1),
        "AdaBoost": AdaBoostClassifier(n_estimators=300, learning_rate=0.1,
                                        random_state=random_state),
        "Gradient Boosting": GradientBoostingClassifier(n_estimators=200, learning_rate=0.05,
                                                         random_state=random_state),
        "XGBoost": XGBClassifier(n_estimators=300, learning_rate=0.05, max_depth=5,
                                  subsample=0.8, colsample_bytree=0.8,
                                  objective="multi:softprob", num_class=3,
                                  random_state=random_state, eval_metric="mlogloss"),
        "Stacking Ensemble": StackingClassifier(
            estimators=[
                ("rf", RandomForestClassifier(n_estimators=300, class_weight="balanced",
                                               random_state=random_state, n_jobs=1)),
                ("et", ExtraTreesClassifier(n_estimators=300, class_weight="balanced",
                                             random_state=random_state, n_jobs=1)),
                ("xgb", XGBClassifier(n_estimators=300, learning_rate=0.05, max_depth=5,
                                       subsample=0.8, colsample_bytree=0.8,
                                       objective="multi:softprob", num_class=3,
                                       random_state=random_state, eval_metric="mlogloss", n_jobs=1)),
            ],
            final_estimator=LogisticRegression(max_iter=1000),
            cv=3, n_jobs=1,
        ),
    }


def get_param_distributions():
    """RandomizedSearchCV spaces using sklearn nested estimator names."""
    return {
        "Random Forest": {
            "clf__n_estimators": [200, 300, 400, 500],
            "clf__max_depth": [None, 10, 20, 30],
            "clf__min_samples_split": [2, 5, 10],
            "clf__max_features": ["sqrt", "log2"],
        },
        "Extra Trees": {
            "clf__n_estimators": [200, 300, 400, 500],
            "clf__max_depth": [None, 10, 20, 30],
            "clf__min_samples_split": [2, 5, 10],
            "clf__max_features": ["sqrt", "log2"],
        },
        "AdaBoost": {
            "clf__n_estimators": [100, 200, 300, 400],
            "clf__learning_rate": [0.01, 0.05, 0.1, 0.5, 1.0],
        },
        "Gradient Boosting": {
            "clf__n_estimators": [100, 200, 300],
            "clf__learning_rate": [0.01, 0.05, 0.1, 0.2],
            "clf__max_depth": [2, 3, 4, 5],
            "clf__subsample": [0.7, 0.8, 0.9, 1.0],
        },
        "XGBoost": {
            "clf__n_estimators": [200, 300, 400],
            "clf__learning_rate": [0.01, 0.05, 0.1],
            "clf__max_depth": [3, 4, 5, 6, 7],
            "clf__subsample": [0.7, 0.8, 0.9, 1.0],
            "clf__colsample_bytree": [0.7, 0.8, 0.9, 1.0],
        },
        "Stacking Ensemble": {
            "clf__rf__max_depth": [10, 20, None],
            "clf__et__max_depth": [10, 20, None],
            "clf__xgb__max_depth": [3, 5],
            "clf__final_estimator__C": [0.1, 1.0, 10.0],
        },
    }


# ----------------------------------------------------------------------
# 3. LEAKAGE-SAFE PIPELINE (scaler -> [feature selection] -> [resampling] -> model)
# ----------------------------------------------------------------------
def build_pipeline(model, resampling, feature_selection, k_features, random_state=42):
    """Build an imblearn Pipeline so that scaling, feature selection, and
    resampling are all re-fit on the training fold only (never the
    validation fold or test set) whenever .fit() is called on this object.
    imblearn's Pipeline (not sklearn's) is required because SMOTE/SMOTETomek
    change the number of rows, which sklearn's own Pipeline does not support.
    """
    steps = [("scaler", StandardScaler())]

    if feature_selection:
        steps.append(("feature_selection", SelectKBest(score_func=f_classif, k=k_features)))

    if resampling == "SMOTE":
        steps.append(("resampler", SMOTE(random_state=random_state)))
    elif resampling == "SMOTE+Tomek":
        steps.append(("resampler", SMOTETomek(random_state=random_state)))
    elif resampling == "Original":
        pass  # no resampling step
    else:
        raise ValueError(f"Unknown resampling option: {resampling}")

    steps.append(("clf", clone(model)))
    return ImbPipeline(steps)


# ----------------------------------------------------------------------
# 4. METRICS
# ----------------------------------------------------------------------
def compute_metrics(y_true, y_pred, y_proba=None):
    metrics = {
        "Accuracy": accuracy_score(y_true, y_pred),
        "Balanced_Accuracy": balanced_accuracy_score(y_true, y_pred),
        "Precision_weighted": precision_score(y_true, y_pred, average="weighted", zero_division=0),
        "Recall_weighted": recall_score(y_true, y_pred, average="weighted", zero_division=0),
        "Macro_F1": f1_score(y_true, y_pred, average="macro", zero_division=0),
        "Primary_Selection_Score": primary_selection_score(y_true, y_pred),
        "F1_weighted": f1_score(y_true, y_pred, average="weighted", zero_division=0),
        "Kappa": cohen_kappa_score(y_true, y_pred),
        "MCC": matthews_corrcoef(y_true, y_pred),
    }
    if y_proba is not None:
        try:
            y_true_bin = label_binarize(y_true, classes=[0, 1, 2])
            metrics["AUC_ovr"] = roc_auc_score(y_true_bin, y_proba, multi_class="ovr")
        except ValueError:
            metrics["AUC_ovr"] = np.nan
    else:
        metrics["AUC_ovr"] = np.nan
    return metrics


ALL_METRICS = ["Primary_Selection_Score", "Accuracy", "Balanced_Accuracy", "Precision_weighted", "Recall_weighted",
               "Macro_F1", "F1_weighted", "Kappa", "MCC", "AUC_ovr"]


CV_CHECKPOINT_PATH = "results/cv_checkpoint.json"


def _json_default(o):
    return o.item() if isinstance(o, np.generic) else str(o)


def load_cv_checkpoints(paths, settings):
    """Collect finished (model, resampling) arms from earlier runs.

    Checkpoints produced with different CV settings are refused: their folds
    or search budget would differ, which would invalidate the paired tests.
    """
    done = {}
    for path in paths:
        if not os.path.exists(path):
            print(f"No checkpoint at {path} yet - nothing to resume from it.")
            continue
        with open(path) as f:
            ckpt = json.load(f)
        if ckpt["settings"] != settings:
            sys.exit(f"Checkpoint {path} was produced with different settings:\n"
                     f"  checkpoint: {ckpt['settings']}\n  this run:   {settings}")
        for arm in ckpt["arms"]:
            done[(arm["Model"], arm["Resampling"])] = arm
    return done


def save_cv_checkpoint(settings, fold_scores, best_params_log):
    """Write every finished arm so an interrupted run loses no completed work."""
    arms = [{"Model": k[0], "Resampling": k[1],
             "fold_scores": {m: v.tolist() for m, v in fold_scores[k].items()},
             "best_params": best_params_log[k]} for k in fold_scores]
    tmp = CV_CHECKPOINT_PATH + ".tmp"
    with open(tmp, "w") as f:
        json.dump({"settings": settings, "arms": arms}, f, indent=2, default=_json_default)
    os.replace(tmp, CV_CHECKPOINT_PATH)


# ----------------------------------------------------------------------
# 5. NESTED REPEATED-CV EVALUATION (outer = evaluation, inner = tuning)
# ----------------------------------------------------------------------
def run_nested_cv(X_train, y_train, models, resampling_list, feature_selection, k_features,
                   n_splits, n_repeats, inner_splits, tune, n_iter_search,
                   random_state=42, resume_from=()):
    """Runs Repeated Stratified K-Fold CV on the 80% training set for every
    (model, resampling) combination, using the SAME outer fold indices for
    every combination (required for valid paired statistical comparison
    later). Hyperparameter tuning, when enabled, happens via an inner
    RandomizedSearchCV fit ONLY on each outer training fold.

    Returns:
      fold_scores: dict[(model_name, resampling)][metric] -> np.array of
                   per-outer-fold scores (length n_splits * n_repeats)
      best_params_log: dict[(model_name, resampling)] -> list of best_params
                        dicts found per outer fold (for transparency/audit)
    """
    outer_cv = RepeatedStratifiedKFold(n_splits=n_splits, n_repeats=n_repeats,
                                        random_state=random_state)
    # Fix the fold split ONCE so every model/resampling arm sees identical
    # train/validation partitions -> makes the later paired test valid.
    fold_indices = list(outer_cv.split(X_train, y_train))
    n_folds = len(fold_indices)
    param_dists = get_param_distributions()

    settings = {"n_splits": n_splits, "n_repeats": n_repeats, "inner_splits": inner_splits,
                "tune": bool(tune), "n_iter_search": n_iter_search, "random_state": random_state,
                "feature_selection": bool(feature_selection), "k_features": k_features,
                "n_train": len(X_train)}
    resumed = load_cv_checkpoints(resume_from, settings)

    fold_scores = {}
    best_params_log = {}

    total_combos = len(models) * len(resampling_list)
    combo_i = 0

    for model_name, model in models.items():
        for resampling in resampling_list:
            combo_i += 1
            key = (model_name, resampling)
            if key in resumed:
                fold_scores[key] = {m: np.array(resumed[key]["fold_scores"][m], dtype=float)
                                    for m in ALL_METRICS}
                best_params_log[key] = resumed[key]["best_params"]
                print(f"\n[{combo_i}/{total_combos}] {model_name} | {resampling} "
                      f"| loaded from checkpoint")
                continue
            fold_scores[key] = {m: np.zeros(n_folds) for m in ALL_METRICS}
            best_params_log[key] = []

            print(f"\n[{combo_i}/{total_combos}] {model_name} | {resampling} "
                  f"| outer folds: {n_folds} | tuning: {tune and model_name in param_dists}")
            t0 = time.time()

            for fi, (tr_idx, val_idx) in enumerate(fold_indices):
                X_tr, X_val = X_train.iloc[tr_idx], X_train.iloc[val_idx]
                y_tr, y_val = y_train.iloc[tr_idx], y_train.iloc[val_idx]

                pipeline = build_pipeline(model, resampling, feature_selection,
                                           k_features, random_state)

                if tune and model_name in param_dists:
                    inner_cv = StratifiedKFold(n_splits=inner_splits, shuffle=True,
                                                random_state=random_state)
                    search = RandomizedSearchCV(
                        pipeline, param_distributions=param_dists[model_name],
                        n_iter=n_iter_search, scoring=PRIMARY_SCORER,
                        cv=inner_cv, random_state=random_state,
                        # Stacking and its base models are single-threaded, so
                        # the search can use every core here as well.
                        n_jobs=-1,
                        refit=True, error_score="raise",
                    )
                    search.fit(X_tr, y_tr)
                    fitted = search.best_estimator_
                    best_params_log[key].append(search.best_params_)
                else:
                    fitted = pipeline
                    fitted.fit(X_tr, y_tr)
                    best_params_log[key].append(None)

                y_pred = fitted.predict(X_val)
                y_proba = fitted.predict_proba(X_val) if hasattr(fitted, "predict_proba") else None
                m = compute_metrics(y_val, y_pred, y_proba)
                for metric_name, val in m.items():
                    fold_scores[key][metric_name][fi] = val

            elapsed = time.time() - t0
            bal_acc_mean = fold_scores[key]["Balanced_Accuracy"].mean()
            f1_mean = fold_scores[key]["Macro_F1"].mean()
            print(f"    -> Balanced Acc = {bal_acc_mean:.4f}, Macro-F1 = {f1_mean:.4f}  ({elapsed:.1f}s)")
            save_cv_checkpoint(settings, fold_scores, best_params_log)

    # Export fold-level observations and tuning choices for reproducibility.
    fold_rows = []
    for (model_name, arm), metric_arrays in fold_scores.items():
        for fold_id in range(n_folds):
            record = {"Model": model_name, "Resampling": arm, "Fold": fold_id + 1}
            record.update({metric: float(vals[fold_id]) for metric, vals in metric_arrays.items()})
            fold_rows.append(record)
    pd.DataFrame(fold_rows).to_csv("results/fold_level_results.csv", index=False)
    with open("results/tuning_history.json", "w") as f:
        json.dump({f"{k[0]} | {k[1]}": v for k, v in best_params_log.items()},
                  f, indent=2, default=str)
    return fold_scores, best_params_log, n_folds, n_splits


# ----------------------------------------------------------------------
# 6. SUMMARY STATISTICS (mean / SD / 95% CI from repeated-CV fold scores)
# ----------------------------------------------------------------------
def summarize_fold_scores(fold_scores):
    rows = []
    for (model_name, resampling), metrics_dict in fold_scores.items():
        row = {"Model": model_name, "Resampling": resampling}
        for metric_name, scores in metrics_dict.items():
            n = len(scores)
            mean = np.mean(scores)
            sd = np.std(scores, ddof=1)
            se = sd / np.sqrt(n)
            tcrit = stats.t.ppf(0.975, df=n - 1)
            ci_low, ci_high = mean - tcrit * se, mean + tcrit * se
            row[f"{metric_name}_mean"] = mean
            row[f"{metric_name}_sd"] = sd
            row[f"{metric_name}_CI95_low"] = ci_low
            row[f"{metric_name}_CI95_high"] = ci_high
        rows.append(row)
    return pd.DataFrame(rows)


def print_comparison_table(summary_df, n_folds, primary_metrics=PRIMARY_METRICS,
                            secondary_metrics=("Accuracy", "AUC_ovr")):
    """Prints every model x resampling combination's cross-validated
    performance (mean, 95% CI) to the console as one table, sorted by the
    pre-registered combined primary-metric score (descending) so the
    selected final model appears first.
    """
    df = summary_df.copy()
    score_cols = [f"{m}_mean" for m in primary_metrics]
    df["selection_score"] = df[score_cols].mean(axis=1)
    df = df.sort_values("selection_score", ascending=False)

    metrics_to_show = list(primary_metrics) + list(secondary_metrics)
    col_w = 24
    header = f"{'Model':<20}{'Resampling':<14}" + "".join(
        f"{m:<{col_w}}" for m in metrics_to_show)

    print("\n" + "=" * len(header))
    print(f"CROSS-VALIDATION MODEL COMPARISON (mean, 95% CI; {n_folds} outer folds)")
    print(f"Sorted by mean of pre-registered primary metrics: {list(primary_metrics)}")
    print("=" * len(header))
    print(header)
    print("-" * len(header))
    for _, row in df.iterrows():
        line = f"{row['Model']:<20}{row['Resampling']:<14}"
        for m in metrics_to_show:
            mean = row[f"{m}_mean"]
            lo = row[f"{m}_CI95_low"]
            hi = row[f"{m}_CI95_high"]
            line += f"{mean:.3f} [{lo:.3f},{hi:.3f}]".ljust(col_w)
        print(line)
    print("=" * len(header))


# ----------------------------------------------------------------------
# 7. PAIRED STATISTICAL COMPARISON (corrected resampled t-test + Wilcoxon)
# ----------------------------------------------------------------------
def corrected_resampled_ttest(diffs, n_train, n_test):
    """Nadeau & Bengio (2003) corrected paired t-test for k-fold /
    repeated k-fold CV. A plain paired t-test on CV fold scores understates
    variance (folds share overlapping training data) and inflates false
    positives; this correction adjusts for that using the train/validation
    size ratio.
    """
    n = len(diffs)
    mean_diff = np.mean(diffs)
    var_diff = np.var(diffs, ddof=1)
    if var_diff == 0:
        return 0.0, 1.0
    correction = (1.0 / n) + (n_test / n_train)
    t_stat = mean_diff / np.sqrt(correction * var_diff)
    df = n - 1
    p_value = 2 * (1 - stats.t.cdf(abs(t_stat), df=df))
    return t_stat, p_value


def compare_models_pairwise(fold_scores, model_resampling_pairs, metric, n_train, n_test, alpha=0.05):
    """Pairwise comparison across a list of (model_name, resampling) keys,
    using the corrected resampled t-test AND Wilcoxon signed-rank test on
    the SAME outer CV folds for both entries in each pair.
    """
    rows = []
    for i in range(len(model_resampling_pairs)):
        for j in range(i + 1, len(model_resampling_pairs)):
            key_a, key_b = model_resampling_pairs[i], model_resampling_pairs[j]
            scores_a = fold_scores[key_a][metric]
            scores_b = fold_scores[key_b][metric]
            diffs = scores_a - scores_b
            t_stat, p_ttest = corrected_resampled_ttest(diffs, n_train, n_test)
            try:
                w_stat, p_wilcoxon = stats.wilcoxon(scores_a, scores_b)
            except ValueError:
                w_stat, p_wilcoxon = np.nan, np.nan
            rows.append({
                "Metric": metric,
                "A": f"{key_a[0]} ({key_a[1]})",
                "B": f"{key_b[0]} ({key_b[1]})",
                "Mean_A": scores_a.mean(),
                "Mean_B": scores_b.mean(),
                "Mean_diff (A-B)": diffs.mean(),
                "corrected_t": t_stat,
                "p_corrected_ttest": p_ttest,
                "wilcoxon_stat": w_stat,
                "p_wilcoxon": p_wilcoxon,
                "significant_at_0.05": (p_ttest < alpha),
            })
    df = pd.DataFrame(rows)
    return df


def apply_holm(df, alpha=0.05):
    """Adjust both statistical tests separately within the declared family."""
    df = df.copy()
    for source, target in [("p_corrected_ttest", "p_holm_corrected_ttest"),
                           ("p_wilcoxon", "p_holm_wilcoxon")]:
        mask = df[source].notna()
        df[target] = np.nan
        if mask.any():
            df.loc[mask, target] = multipletests(
                df.loc[mask, source].to_numpy(), alpha=alpha, method="holm"
            )[1]
    df["significant_at_0.05"] = df["p_holm_corrected_ttest"] < alpha
    df["significant_holm_wilcoxon"] = df["p_holm_wilcoxon"] < alpha
    return df


# ----------------------------------------------------------------------
# 8. MODEL SELECTION (pre-defined primary metrics, item 4/5)
# ----------------------------------------------------------------------
def select_final_model(summary_df, primary_metrics=PRIMARY_METRICS):
    """Rank by the AVERAGE of the pre-registered primary metrics' means.
    This decision rule is fixed BEFORE looking at results (declared at the
    top of this file) - it does not change based on which model happens to
    win on Accuracy or AUC.
    """
    score_cols = [f"{m}_mean" for m in primary_metrics]
    summary_df = summary_df.copy()
    summary_df["selection_score"] = summary_df[score_cols].mean(axis=1)
    best_row = summary_df.sort_values(
        ["selection_score", "Balanced_Accuracy_mean", "Macro_F1_mean", "Model", "Resampling"],
        ascending=[False, False, False, True, True]).iloc[0]
    return best_row["Model"], best_row["Resampling"], summary_df


# ----------------------------------------------------------------------
# 9. FINAL REFIT ON FULL 80% TRAIN + TUNING (still no test-set contact)
# ----------------------------------------------------------------------
def tune_and_fit_final(model_name, resampling, X_train, y_train, feature_selection,
                        k_features, cv_splits, n_iter_search, tune, random_state=42):
    model = get_base_models(random_state)[model_name]
    pipeline = build_pipeline(model, resampling, feature_selection, k_features, random_state)
    param_dists = get_param_distributions()

    if tune and model_name in param_dists:
        cv = StratifiedKFold(n_splits=cv_splits, shuffle=True, random_state=random_state)
        search = RandomizedSearchCV(
            pipeline, param_distributions=param_dists[model_name],
            n_iter=n_iter_search, scoring=PRIMARY_SCORER,
            cv=cv, random_state=random_state,
            n_jobs=-1,
            refit=True, error_score="raise",
        )
        search.fit(X_train, y_train)
        print(f"Final tuned hyperparameters for {model_name}: {search.best_params_}")
        return search.best_estimator_, search.best_params_
    else:
        pipeline.fit(X_train, y_train)
        return pipeline, None


# ----------------------------------------------------------------------
# 10. FINAL, SINGLE EVALUATION ON THE LOCKED 20% TEST SET + BOOTSTRAP CI
# ----------------------------------------------------------------------
def bootstrap_ci_metrics(y_true, y_pred, y_proba, n_boot=2000, random_state=42, alpha=0.05):
    rng = np.random.RandomState(random_state)
    y_true = np.asarray(y_true)
    y_pred = np.asarray(y_pred)
    n = len(y_true)
    boot_metrics = {m: [] for m in ALL_METRICS}

    for _ in range(n_boot):
        idx = rng.randint(0, n, n)
        yt, yp = y_true[idx], y_pred[idx]
        yprob = y_proba[idx] if y_proba is not None else None
        if len(np.unique(yt)) < 2:
            continue
        m = compute_metrics(yt, yp, yprob)
        for k, v in m.items():
            boot_metrics[k].append(v)

    ci = {}
    for k, vals in boot_metrics.items():
        vals = np.array(vals)
        vals = vals[~np.isnan(vals)]
        if len(vals) == 0:
            ci[k] = (np.nan, np.nan)
        else:
            ci[k] = (np.percentile(vals, 100 * alpha / 2), np.percentile(vals, 100 * (1 - alpha / 2)))
    return ci


def evaluate_on_locked_test(fitted_pipeline, X_test, y_test, n_boot=2000, random_state=42):
    y_pred = fitted_pipeline.predict(X_test)
    y_proba = fitted_pipeline.predict_proba(X_test) if hasattr(fitted_pipeline, "predict_proba") else None
    point_metrics = compute_metrics(y_test, y_pred, y_proba)
    ci = bootstrap_ci_metrics(y_test, y_pred, y_proba, n_boot=n_boot, random_state=random_state)
    return point_metrics, ci, y_pred, y_proba


# ----------------------------------------------------------------------
# 11. PLOTS
# ----------------------------------------------------------------------
def save_plots(model_name, y_test, y_pred, fitted_pipeline, feature_names,
               feature_selection, point_metrics):
    fig, ax = plt.subplots(figsize=(6, 5))
    cm = confusion_matrix(y_test, y_pred)
    sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", ax=ax,
                xticklabels=CLASS_NAMES, yticklabels=CLASS_NAMES)
    ax.set_title(f"Confusion Matrix (held-out test) - {model_name}\n"
                 f"Balanced Acc: {point_metrics['Balanced_Accuracy']:.2%}")
    plt.tight_layout()
    plt.savefig("plots/confusion_matrix.png", dpi=150)
    plt.close()

    try:
        clf = fitted_pipeline.named_steps["clf"]
    except Exception:
        clf = None

    if clf is not None and hasattr(clf, "feature_importances_"):
        if feature_selection and "feature_selection" in fitted_pipeline.named_steps:
            selector = fitted_pipeline.named_steps["feature_selection"]
            selected_idx = selector.get_support(indices=True)
            feat_names = [feature_names[i] for i in selected_idx]
        else:
            feat_names = feature_names

        imp = clf.feature_importances_
        if len(imp) == len(feat_names):
            feat_df = pd.DataFrame({"feature": feat_names, "importance": imp})
            feat_df = feat_df.sort_values("importance", ascending=True).tail(15)
            fig, ax = plt.subplots(figsize=(8, 6))
            colors = plt.cm.YlOrRd(np.linspace(0.3, 0.9, len(feat_df)))
            ax.barh(feat_df["feature"], feat_df["importance"], color=colors)
            ax.set_xlabel("Importance")
            ax.set_title(f"Top {len(feat_df)} Features - {model_name} (final model)")
            ax.grid(True, alpha=0.3, axis="x")
            plt.tight_layout()
            plt.savefig("plots/feature_importance.png", dpi=150)
            plt.close()


# ----------------------------------------------------------------------
# 12. SAVE ARTIFACTS
# ----------------------------------------------------------------------
def save_cv_summary_csv(summary_df):
    filename = f"results/cv_summary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    summary_df.to_csv(filename, index=False)
    print(f"CV summary (mean/SD/95% CI per model x resampling) saved to {filename}")
    return filename


def save_comparison_csv(comparison_df, name):
    filename = f"results/{name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    comparison_df.to_csv(filename, index=False)
    print(f"{name} saved to {filename}")
    return filename


def save_final_test_results(final_model_name, final_resampling, point_metrics, ci,
                             best_params):
    result = {
        "final_model": final_model_name,
        "final_resampling": final_resampling,
        "best_hyperparameters": best_params,
        "test_point_estimates": point_metrics,
        "test_bootstrap_95CI": {k: list(v) for k, v in ci.items()},
    }
    filename = f"results/final_test_results_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    with open(filename, "w") as f:
        json.dump(result, f, indent=2, default=float)
    print(f"Final locked-test results saved to {filename}")
    return filename


def save_model_bundle(fitted_pipeline, final_model_name, final_resampling, best_params,
                       feature_names):
    bundle = {
        "pipeline": fitted_pipeline,
        "model_name": final_model_name,
        "resampling": final_resampling,
        "best_hyperparameters": best_params,
        "feature_names": feature_names,
        "class_labels": CLASS_NAMES,
    }
    safe_name = final_model_name.replace(" ", "_")
    filename = f"models/{safe_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pkl"
    joblib.dump(bundle, filename)
    print(f"Final model bundle saved to {filename}")
    return filename


def generate_word_report(cv_summary_df, model_comparison_df, resampling_comparison_df,
                          final_model_name, final_resampling, point_metrics, ci,
                          best_params, args):
    if not DOCX_AVAILABLE:
        print("python-docx not installed. Word report skipped.")
        return None

    doc = Document()
    title = doc.add_heading("CTG ML Pipeline Report - Nested CV / Leakage-Safe Version", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Outer CV: Repeated Stratified {args.n_splits}-Fold x {args.n_repeats} repeats "
                       f"({args.n_splits * args.n_repeats} total outer folds) on the 80% training set")
    doc.add_paragraph(f"Feature selection: {'Enabled (top ' + str(args.k_features) + ')' if args.feature_selection else 'Disabled'}")
    doc.add_paragraph(f"Hyperparameter tuning: {'RandomizedSearchCV (' + str(args.n_iter_search) + ' iters, inner ' + str(args.inner_splits) + '-fold CV)' if args.tune else 'Disabled (default hyperparameters)'}")
    doc.add_paragraph(f"Pre-registered primary metrics: {', '.join(PRIMARY_METRICS)}")

    doc.add_heading("Methodology note on data leakage", level=1)
    doc.add_paragraph(METHODOLOGY_LEAKAGE_STATEMENT)

    doc.add_heading("Cross-Validation Summary (mean, SD, 95% CI per model x resampling)", level=1)
    display_cols = ["Model", "Resampling"] + [f"{m}_mean" for m in PRIMARY_METRICS + SECONDARY_METRICS]
    table = doc.add_table(rows=len(cv_summary_df) + 1, cols=len(display_cols))
    table.style = "Table Grid"
    for i, h in enumerate(display_cols):
        table.rows[0].cells[i].text = h.replace("_mean", "")
    for i, (_, r) in enumerate(cv_summary_df.iterrows()):
        row = table.rows[i + 1]
        row.cells[0].text = str(r["Model"])
        row.cells[1].text = str(r["Resampling"])
        for j, col in enumerate(display_cols[2:], start=2):
            val = r[col]
            row.cells[j].text = f"{val:.4f}" if pd.notna(val) else "NA"

    doc.add_heading("Pairwise Model Comparison (corrected resampled t-test + Wilcoxon)", level=1)
    doc.add_paragraph("Comparisons use the same CV folds for every model (paired design). "
                       "A plain paired t-test on CV folds is known to understate variance; "
                       "the corrected resampled t-test (Nadeau & Bengio, 2003) adjusts for this.")
    if model_comparison_df is not None and len(model_comparison_df) > 0:
        cols = ["A", "B", "Mean_A", "Mean_B", "Mean_diff (A-B)", "p_corrected_ttest", "p_holm_corrected_ttest", "p_wilcoxon", "p_holm_wilcoxon", "significant_at_0.05"]
        table = doc.add_table(rows=len(model_comparison_df) + 1, cols=len(cols))
        table.style = "Table Grid"
        for i, h in enumerate(cols):
            table.rows[0].cells[i].text = h
        for i, (_, r) in enumerate(model_comparison_df.iterrows()):
            row = table.rows[i + 1]
            for j, col in enumerate(cols):
                val = r[col]
                row.cells[j].text = f"{val:.4f}" if isinstance(val, float) else str(val)

    doc.add_heading("Resampling Comparison per Model (Original vs SMOTE vs SMOTE+Tomek)", level=1)
    if resampling_comparison_df is not None and len(resampling_comparison_df) > 0:
        cols = list(resampling_comparison_df.columns)
        table = doc.add_table(rows=len(resampling_comparison_df) + 1, cols=len(cols))
        table.style = "Table Grid"
        for i, h in enumerate(cols):
            table.rows[0].cells[i].text = h
        for i, (_, r) in enumerate(resampling_comparison_df.iterrows()):
            row = table.rows[i + 1]
            for j, col in enumerate(cols):
                val = r[col]
                row.cells[j].text = f"{val:.4f}" if isinstance(val, float) else str(val)

    doc.add_heading("Final Model Selection", level=1)
    doc.add_paragraph(
        f"Model selection was based primarily on the predefined Balanced Accuracy and "
        f"Macro-F1 metrics because of the class imbalance in the dataset. Accuracy, "
        f"Precision, Recall, and ROC-AUC were additionally reported to provide a "
        f"comprehensive assessment. Model comparisons were performed using the same "
        f"cross-validation folds, and the final selected model "
        f"({final_model_name}, {final_resampling} resampling) was subsequently "
        f"evaluated once on the held-out test set."
    )
    if best_params:
        doc.add_paragraph(f"Final tuned hyperparameters: {best_params}")

    doc.add_heading("Final Held-Out Test Results (evaluated once)", level=1)
    cols = ["Metric", "Point estimate", "Bootstrap 95% CI"]
    metrics_to_show = PRIMARY_METRICS + SECONDARY_METRICS
    table = doc.add_table(rows=len(metrics_to_show) + 1, cols=len(cols))
    table.style = "Table Grid"
    for i, h in enumerate(cols):
        table.rows[0].cells[i].text = h
    for i, m in enumerate(metrics_to_show):
        row = table.rows[i + 1]
        row.cells[0].text = m
        row.cells[1].text = f"{point_metrics.get(m, np.nan):.4f}"
        lo, hi = ci.get(m, (np.nan, np.nan))
        row.cells[2].text = f"[{lo:.4f}, {hi:.4f}]"

    if os.path.exists("plots/confusion_matrix.png"):
        doc.add_heading("Confusion Matrix (final model, held-out test)", level=1)
        doc.add_picture("plots/confusion_matrix.png", width=Inches(5))

    doc.add_heading("Data Availability", level=1)
    doc.add_paragraph(DATA_AVAILABILITY_STATEMENT)

    doc.add_heading("Limitations", level=1)
    doc.add_paragraph(LIMITATIONS_STATEMENT)

    filename = f"results/report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    doc.save(filename)
    print(f"Word report saved to {filename}")
    return filename


# ----------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------
def main():
    parser = argparse.ArgumentParser(description="CTG ML Pipeline v1 (22-9) - Nested CV / Leakage-Safe")
    parser.add_argument("--data", type=str, default="CTG_dataset.csv")
    parser.add_argument("--models", nargs="+",
                        default=["Random Forest", "Extra Trees", "AdaBoost",
                                 "Gradient Boosting", "XGBoost", "Stacking Ensemble"])
    parser.add_argument("--resampling", nargs="+", default=["Original", "SMOTE", "SMOTE+Tomek"],
                        choices=["Original", "SMOTE", "SMOTE+Tomek"])
    parser.add_argument("--feature_selection", action="store_true")
    parser.add_argument("--k_features", type=int, default=15)
    parser.add_argument("--n_splits", type=int, default=5, help="Outer CV folds")
    parser.add_argument("--n_repeats", type=int, default=3, help="Outer CV repeats")
    parser.add_argument("--inner_splits", type=int, default=3, help="Inner CV folds for tuning")
    parser.add_argument("--tune", action=argparse.BooleanOptionalAction, default=True,
                        help="Enable nested RandomizedSearchCV tuning (default: enabled)")
    parser.add_argument("--n_iter_search", type=int, default=30)
    parser.add_argument("--n_boot", type=int, default=2000, help="Bootstrap resamples for final test CI")
    parser.add_argument("--report", action="store_true")
    parser.add_argument("--random_state", type=int, default=42)
    parser.add_argument("--resume_from", nargs="*", default=[],
                        help="cv_checkpoint.json files from earlier runs; their finished "
                             "model/resampling arms are reused instead of recomputed")
    parser.add_argument("--cv_only", action="store_true",
                        help="Stop after nested CV (for partial runs split by --models). "
                             "No model is selected and the locked test set is not touched.")
    args = parser.parse_args()

    if args.tune and "Stacking Ensemble" in args.models:
        print("NOTICE: Stacking Ensemble is tuned with a restricted parameter space "
              "and 3-fold internal stacking CV; its training time may be substantial.")

    rs = args.random_state

    # Step 1: locked split
    X_train, X_test, y_train, y_test, feature_names = load_and_split_data(
        args.data, test_size=0.20, random_state=rs
    )

    all_models = get_base_models(rs)
    models = {name: all_models[name] for name in args.models if name in all_models}
    if not models:
        print("No valid models selected. Exiting.")
        sys.exit(1)

    # Step 2: nested repeated CV on the 80% training set only
    fold_scores, best_params_log, n_folds, n_splits = run_nested_cv(
        X_train, y_train, models, args.resampling, args.feature_selection, args.k_features,
        args.n_splits, args.n_repeats, args.inner_splits, args.tune, args.n_iter_search, rs,
        resume_from=args.resume_from,
    )
    if args.cv_only:
        print(f"\nCV-only run finished. Checkpoint saved to {CV_CHECKPOINT_PATH}; "
              "the locked test set was not used.")
        return

    # Step 3: summary stats (mean/SD/95% CI)
    summary_df = summarize_fold_scores(fold_scores)
    save_cv_summary_csv(summary_df)
    print_comparison_table(summary_df, n_folds)

    # One predefined score controls both resampling choice and final model.
    summary_df["selection_score"] = summary_df[
        ["Balanced_Accuracy_mean", "Macro_F1_mean"]].mean(axis=1)
    best_per_model = (summary_df.sort_values(
        ["selection_score", "Balanced_Accuracy_mean", "Macro_F1_mean", "Model", "Resampling"],
        ascending=[False, False, False, True, True]
    ).groupby("Model").head(1))
    # Step 4: pairwise model comparison using best arm by the SAME rule
    model_pairs = list(zip(best_per_model["Model"], best_per_model["Resampling"]))
    model_comparison_df = compare_models_pairwise(
        fold_scores, model_pairs, metric=PRIMARY_COMPOSITE,
        n_train=len(X_train) * (n_splits - 1) / n_splits, n_test=len(X_train) / n_splits,
    )
    if not model_comparison_df.empty:
        model_comparison_df = apply_holm(model_comparison_df)
    save_comparison_csv(model_comparison_df, "model_comparison")

    # Step 5: resampling comparison, per model (Original vs SMOTE vs SMOTE+Tomek)
    resampling_rows = []
    for model_name in models:
        arms = [(model_name, r) for r in args.resampling]
        if len(arms) > 1:
            cmp_df = compare_models_pairwise(
                fold_scores, arms, metric=PRIMARY_COMPOSITE,
                n_train=len(X_train) * (n_splits - 1) / n_splits, n_test=len(X_train) / n_splits,
            )
            resampling_rows.append(cmp_df)
    resampling_comparison_df = pd.concat(resampling_rows, ignore_index=True) if resampling_rows else pd.DataFrame()
    # Treat all resampling comparisons as one predefined Holm family.
    if not resampling_comparison_df.empty:
        resampling_comparison_df = apply_holm(resampling_comparison_df)
    save_comparison_csv(resampling_comparison_df, "resampling_comparison")

    # Step 6: model selection using PRE-REGISTERED primary metrics
    final_model_name, final_resampling, summary_df = select_final_model(summary_df)
    print(f"\n{'='*70}\nSELECTED FINAL MODEL: {final_model_name} | Resampling: {final_resampling}")
    print(f"(selected by mean of pre-registered primary metrics: {PRIMARY_METRICS})\n{'='*70}")

    # Step 7: final tuning + refit on the FULL 80% training set (test still untouched)
    final_pipeline, best_params = tune_and_fit_final(
        final_model_name, final_resampling, X_train, y_train, args.feature_selection,
        args.k_features, args.n_splits, args.n_iter_search, args.tune, rs,
    )

    # Step 8: the ONE and ONLY evaluation on the locked 20% test set
    point_metrics, ci, y_pred, y_proba = evaluate_on_locked_test(
        final_pipeline, X_test, y_test, n_boot=args.n_boot, random_state=rs
    )
    print("\nFinal held-out test performance (evaluated once):")
    for m in PRIMARY_METRICS + SECONDARY_METRICS:
        lo, hi = ci.get(m, (np.nan, np.nan))
        print(f"  {m}: {point_metrics.get(m, np.nan):.4f}   95% Bootstrap CI [{lo:.4f}, {hi:.4f}]")

    print("\nClassification report (held-out test):")
    print(classification_report(y_test, y_pred, target_names=CLASS_NAMES, zero_division=0))

    save_final_test_results(final_model_name, final_resampling, point_metrics, ci, best_params)
    save_plots(final_model_name, y_test, y_pred, final_pipeline, feature_names,
               args.feature_selection, point_metrics)
    save_model_bundle(final_pipeline, final_model_name, final_resampling, best_params, feature_names)

    if args.report:
        generate_word_report(summary_df, model_comparison_df, resampling_comparison_df,
                              final_model_name, final_resampling, point_metrics, ci,
                              best_params, args)

    print("\nPipeline completed.")


if __name__ == "__main__":
    main()
