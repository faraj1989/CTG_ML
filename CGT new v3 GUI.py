"""
CTG ML Pipeline GUI – Full Pipeline with No Leakage & Advanced Enhancements
===========================================================================
- Loads CTG_dataset.csv (semicolon)
- Removes duplicates
- Converts NSP to 0,1,2
- Splits 80/20 stratified
- User selects preprocessing: Original, SMOTE, SMOTE+Tomek
- NEW: Optional Feature Selection (SelectKBest)
- Trains: Random Forest, Extra Trees, AdaBoost, Gradient Boosting, XGBoost, Stacking Ensemble
- Displays metrics, confusion matrix, classification report, feature importance
- Generates Word report
- Full console-like logging in the GUI log window
- Dataset Info tab with comprehensive dataset statistics
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, scrolledtext
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
import seaborn as sns
import warnings
import time
import os
from threading import Thread
from datetime import datetime
from io import BytesIO

from sklearn.model_selection import train_test_split, StratifiedKFold, cross_val_score
from sklearn.preprocessing import StandardScaler, label_binarize
from sklearn.feature_selection import SelectKBest, f_classif
from sklearn.ensemble import (RandomForestClassifier, ExtraTreesClassifier, AdaBoostClassifier,
                              GradientBoostingClassifier, StackingClassifier)
from sklearn.linear_model import LogisticRegression
from xgboost import XGBClassifier
from sklearn.metrics import (accuracy_score, precision_score, recall_score, f1_score,
                             balanced_accuracy_score, cohen_kappa_score, matthews_corrcoef,
                             confusion_matrix, classification_report, roc_auc_score, roc_curve)
from sklearn.pipeline import Pipeline
from imblearn.over_sampling import SMOTE
from imblearn.combine import SMOTETomek

# Word report
try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("python-docx not installed. Run: pip install python-docx")

warnings.filterwarnings("ignore")

# Create folders
os.makedirs("cleaned_data", exist_ok=True)
os.makedirs("models", exist_ok=True)
os.makedirs("results", exist_ok=True)


class CTGPipelineGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("CTG ML Pipeline – Advanced GUI")
        self.root.geometry("1500x900")
        self.root.configure(bg="#f5f7fa")

        # Data containers
        self.df = None
        self.X = None
        self.y = None
        self.X_train = None
        self.X_test = None
        self.y_train = None
        self.y_test = None
        self.scaler = StandardScaler()
        self.selector = None
        self.results = []
        self.models = {}
        self.feature_names = None
        self.class_names = ["Normal", "Suspect", "Pathologic"]

        # GUI setup
        self.setup_gui()
        self.load_data()  # load default dataset

    # ----- Data Loading -----
    def load_data(self, path=None):
        try:
            if path is None:
                path = "CTG_dataset.csv"  # adjust if needed

            if not os.path.exists(path):
                self.log(f"Warning: {path} not found. Please load manually via File menu.")
                return

            self.df = pd.read_csv(path, sep=";")
            self.df.columns = self.df.columns.str.strip()
            self.log(f"Loaded {len(self.df)} rows, {len(self.df.columns)} columns")

            # Remove duplicates
            dup_count = self.df.duplicated().sum()
            self.df = self.df.drop_duplicates()
            self.log(f"After removing duplicates: {len(self.df)}")
            if dup_count > 0:
                self.log(f"  (Removed {dup_count} duplicate rows)")

            # Target conversion
            self.X = self.df.drop(columns=["NSP"])
            self.y = self.df["NSP"] - 1
            self.feature_names = self.X.columns.tolist()

            # Split
            self.X_train, self.X_test, self.y_train, self.y_test = train_test_split(
                self.X, self.y, test_size=0.20, random_state=42, stratify=self.y
            )
            self.log(f"Train: {len(self.X_train)}, Test: {len(self.X_test)}")

            # Class distribution
            nsp = self.df["NSP"].value_counts()
            self.log(
                f"Class distribution — Normal: {nsp.get(1, 0)}, Suspect: {nsp.get(2, 0)}, Pathologic: {nsp.get(3, 0)}")

            self.update_dashboard()
            self.populate_feature_entries()
            self.update_dataset_info()
            self.status_var.set("Dataset loaded")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to load data: {e}")

    # ----- GUI Setup -----
    def setup_gui(self):
        menubar = tk.Menu(self.root)
        self.root.config(menu=menubar)
        file_menu = tk.Menu(menubar, tearoff=0)
        menubar.add_cascade(label="File", menu=file_menu)
        file_menu.add_command(label="Load Dataset", command=self.load_data_dialog)
        file_menu.add_command(label="Generate Report", command=self.generate_word_report)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.root.quit)

        main = tk.Frame(self.root, bg="#f5f7fa")
        main.pack(fill="both", expand=True, padx=10, pady=10)

        tk.Label(main, text="CTG ML Pipeline – Advanced GUI", font=("Arial", 18, "bold"),
                 bg="#f5f7fa", fg="#2c3e50").pack(pady=(0, 5))
        tk.Label(main, text="Preprocessing: Original / SMOTE / SMOTE+Tomek | Models: RF, ET, Ada, GBM, XGB, Stacking",
                 font=("Arial", 11), bg="#f5f7fa", fg="#666").pack(pady=(0, 10))

        self.notebook = ttk.Notebook(main)
        self.notebook.pack(fill="both", expand=True)

        self.create_dashboard_tab()
        self.create_dataset_info_tab()
        self.create_training_tab()
        self.create_prediction_tab()
        self.create_results_tab()

        self.status_var = tk.StringVar(value="Ready")
        tk.Label(main, textvariable=self.status_var, font=("Arial", 9),
                 bg="#f5f7fa", fg="#555").pack(side="bottom", fill="x", pady=(5, 0))

    # ----- Dashboard Tab -----
    def create_dashboard_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="📊 Dashboard")

        metrics_frame = tk.Frame(tab, bg="#f5f7fa")
        metrics_frame.pack(fill="x", padx=20, pady=10)

        self.metric_cards = {}
        metrics = [
            ("Total Samples", "0", "#667eea"),
            ("Features", "0", "#3498db"),
            ("Normal", "0", "#2ecc71"),
            ("Suspect", "0", "#f39c12"),
            ("Pathologic", "0", "#e74c3c"),
            ("Best Accuracy", "0%", "#9b59b6"),
            ("Best F1", "0", "#1abc9c"),
            ("Best MCC", "0", "#e67e22")
        ]
        for i, (label, val, color) in enumerate(metrics):
            card = tk.Frame(metrics_frame, bg="white", relief="raised", bd=2)
            card.grid(row=i // 4, column=i % 4, padx=10, pady=10, sticky="nsew")
            tk.Label(card, text=label, font=("Arial", 10), bg="white", fg="#666").pack(pady=(10, 5))
            lbl = tk.Label(card, text=val, font=("Arial", 20, "bold"), bg="white", fg=color)
            lbl.pack(pady=(0, 10))
            self.metric_cards[label] = lbl

        for i in range(4):
            metrics_frame.grid_columnconfigure(i, weight=1)

        perf_frame = tk.LabelFrame(tab, text="Model Performance", font=("Arial", 12, "bold"),
                                   bg="white", fg="#2c3e50")
        perf_frame.pack(fill="both", expand=True, padx=20, pady=10)

        self.dash_fig = Figure(figsize=(8, 4), dpi=100)
        self.dash_ax = self.dash_fig.add_subplot(111)
        self.dash_canvas = FigureCanvasTkAgg(self.dash_fig, master=perf_frame)
        self.dash_canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=10)

    # ----- Dataset Info Tab -----
    def create_dataset_info_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="📋 Dataset Info")

        self.dataset_info_text = scrolledtext.ScrolledText(tab, font=("Courier", 9), wrap="word")
        self.dataset_info_text.pack(fill="both", expand=True, padx=10, pady=10)

        refresh_btn = tk.Button(tab, text="🔄 Refresh Info", command=self.update_dataset_info,
                                bg="#3498db", fg="white", font=("Arial", 10, "bold"))
        refresh_btn.pack(pady=5)

    def update_dataset_info(self):
        if self.df is None:
            self.dataset_info_text.delete(1.0, tk.END)
            self.dataset_info_text.insert(tk.END, "No dataset loaded. Please load a dataset.")
            return

        self.dataset_info_text.delete(1.0, tk.END)
        info = []
        info.append("=" * 70)
        info.append("CTG DATASET INFORMATION")
        info.append("=" * 70)
        info.append(f"Shape: {self.df.shape[0]} rows, {self.df.shape[1]} columns")
        info.append(f"Duplicate rows: {self.df.duplicated().sum()}")
        info.append(f"Missing values: {self.df.isnull().sum().sum()}")
        info.append("")

        info.append("-" * 70)
        info.append("COLUMN NAMES AND DATA TYPES")
        info.append("-" * 70)
        for col in self.df.columns:
            info.append(f"{col:20s} : {str(self.df[col].dtype):10s}  unique={self.df[col].nunique()}")
        info.append("")

        info.append("-" * 70)
        info.append("TARGET DISTRIBUTION (NSP)")
        info.append("-" * 70)
        nsp_counts = self.df["NSP"].value_counts().sort_index()
        total = len(self.df)
        for cls in sorted(nsp_counts.index):
            count = nsp_counts[cls]
            pct = count / total * 100
            class_name = {1: "Normal", 2: "Suspect", 3: "Pathologic"}.get(cls, str(cls))
            info.append(f"Class {cls} ({class_name:10s}) : {count:6d}  ({pct:5.2f}%)")
        info.append("")

        info.append("-" * 70)
        info.append("STATISTICAL SUMMARY (numeric features)")
        info.append("-" * 70)
        numeric_cols = self.df.select_dtypes(include=[np.number]).columns
        summary = self.df[numeric_cols].describe().T
        for col in summary.index:
            row = summary.loc[col]
            info.append(f"{col:20s} : mean={row['mean']:8.2f}  std={row['std']:8.2f}  "
                        f"min={row['min']:8.2f}  25%={row['25%']:8.2f}  50%={row['50%']:8.2f}  75%={row['75%']:8.2f}  max={row['max']:8.2f}")
        info.append("")

        info.append("-" * 70)
        info.append("CORRELATION WITH NSP (top 10 absolute)")
        info.append("-" * 70)
        corr = self.df.corr(numeric_only=True)["NSP"].drop("NSP")
        corr_sorted = corr.abs().sort_values(ascending=False)
        for feat in corr_sorted.index[:10]:
            info.append(f"{feat:20s} : {corr[feat]:+8.4f}  (abs={corr_sorted[feat]:.4f})")
        info.append("")

        info.append("-" * 70)
        info.append("MEMORY USAGE")
        info.append("-" * 70)
        mem = self.df.memory_usage(deep=True).sum() / 1024 ** 2
        info.append(f"Total memory: {mem:.2f} MB")
        info.append("")

        info.append("=" * 70)
        info.append("END OF DATASET INFORMATION")
        info.append("=" * 70)

        self.dataset_info_text.insert(tk.END, "\n".join(info))

    # ----- Training Tab -----
    def create_training_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="🎛️ Training")

        left = tk.Frame(tab, bg="white", relief="raised", bd=2)
        left.pack(side="left", fill="both", expand=True, padx=10, pady=10)

        tk.Label(left, text="Training Configuration", font=("Arial", 16, "bold"), bg="white").pack(pady=10)

        # Preprocessing selection
        tk.Label(left, text="Preprocessing:", font=("Arial", 12), bg="white").pack(pady=(5, 0))
        self.prep_var = tk.StringVar(value="Original")
        prep_combo = ttk.Combobox(left, textvariable=self.prep_var,
                                  values=["Original", "SMOTE", "SMOTE+Tomek"],
                                  state="readonly", width=20)
        prep_combo.pack(pady=5)

        # Feature Selection Toggle
        self.feat_sel_var = tk.BooleanVar(value=False)
        fs_cb = tk.Checkbutton(left, text="Apply Feature Selection (Top 15)", variable=self.feat_sel_var, bg="white")
        fs_cb.pack(pady=5)

        # Model selection (checkboxes)
        tk.Label(left, text="Select Models:", font=("Arial", 12), bg="white").pack(pady=(10, 5))
        model_names = ["Random Forest", "Extra Trees", "AdaBoost", "Gradient Boosting", "XGBoost", "Stacking Ensemble"]
        self.model_vars = {}
        for name in model_names:
            var = tk.BooleanVar(value=True)
            self.model_vars[name] = var
            cb = tk.Checkbutton(left, text=name, variable=var, bg="white")
            cb.pack(anchor="w", padx=30)

        # Train button
        self.train_btn = tk.Button(left, text="🚀 Train Selected Models",
                                   command=self.train_models,
                                   bg="#2ecc71", fg="white", font=("Arial", 14, "bold"),
                                   padx=20, pady=10)
        self.train_btn.pack(pady=20)

        self.progress = ttk.Progressbar(left, mode="indeterminate")
        self.progress.pack(fill="x", padx=20, pady=10)

        # Log area with Clear button
        log_frame = tk.Frame(left, bg="white")
        log_frame.pack(fill="both", expand=True, padx=10, pady=5)
        tk.Button(log_frame, text="Clear Log", command=lambda: self.log_text.delete(1.0, tk.END),
                  bg="#3498db", fg="white", font=("Arial", 9)).pack(anchor="e", pady=(0, 5))
        self.log_text = scrolledtext.ScrolledText(log_frame, height=15, font=("Courier", 9), wrap="word")
        self.log_text.pack(fill="both", expand=True)

        # Right panel: Feature Importance
        right = tk.Frame(tab, bg="white", relief="raised", bd=2)
        right.pack(side="right", fill="both", expand=True, padx=10, pady=10)

        tk.Label(right, text="Feature Importance (Random Forest)", font=("Arial", 12, "bold"), bg="white").pack(pady=5)
        self.fi_fig = Figure(figsize=(6, 5), dpi=100)
        self.fi_ax = self.fi_fig.add_subplot(111)
        self.fi_canvas = FigureCanvasTkAgg(self.fi_fig, master=right)
        self.fi_canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=5)

    # ----- Prediction Tab -----
    def create_prediction_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="🔮 Predictions")

        left = tk.Frame(tab, bg="white", relief="raised", bd=2)
        left.pack(side="left", fill="both", expand=True, padx=10, pady=10)
        tk.Label(left, text="Feature Input", font=("Arial", 16, "bold"), bg="white").pack(pady=10)

        sample_frame = tk.Frame(left, bg="white")
        sample_frame.pack(fill="x", padx=20, pady=10)
        tk.Label(sample_frame, text="Load Sample:", bg="white").pack(side="left", padx=5)
        self.sample_var = tk.StringVar(value="Random")
        combo = ttk.Combobox(sample_frame, textvariable=self.sample_var,
                             values=["Random", "Normal", "Suspect", "Pathologic", "Custom"],
                             state="readonly", width=20)
        combo.pack(side="left", padx=5)
        combo.bind("<<ComboboxSelected>>", self.load_sample)

        canvas = tk.Canvas(left, bg="white")
        scrollbar = tk.Scrollbar(left, orient="vertical", command=canvas.yview)
        self.feature_frame = tk.Frame(canvas, bg="white")
        self.feature_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.feature_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True, padx=20, pady=10)
        scrollbar.pack(side="right", fill="y")
        self.feature_entries = {}

        tk.Label(left, text="Select Model:", font=("Arial", 12), bg="white").pack(pady=(10, 5))
        self.pred_model_var = tk.StringVar(value="Stacking Ensemble")
        self.pred_combo = ttk.Combobox(left, textvariable=self.pred_model_var,
                                       values=["Random Forest", "Extra Trees", "AdaBoost", "Gradient Boosting",
                                               "XGBoost", "Stacking Ensemble"],
                                       state="readonly", width=20)
        self.pred_combo.pack(pady=5)

        self.predict_btn = tk.Button(left, text="🔮 Make Prediction",
                                     command=self.make_prediction,
                                     bg="#9b59b6", fg="white", font=("Arial", 14, "bold"),
                                     padx=20, pady=10)
        self.predict_btn.pack(pady=20)

        right = tk.Frame(tab, bg="white", relief="raised", bd=2)
        right.pack(side="right", fill="both", expand=True, padx=10, pady=10)
        tk.Label(right, text="Prediction Results", font=("Arial", 16, "bold"), bg="white").pack(pady=10)

        self.pred_result_frame = tk.Frame(right, bg="white", height=120)
        self.pred_result_frame.pack(fill="x", padx=20, pady=10)
        self.pred_label = tk.Label(self.pred_result_frame, text="", font=("Arial", 18, "bold"), bg="white")
        self.pred_label.pack(expand=True, fill="both", pady=20)
        self.pred_details = scrolledtext.ScrolledText(right, height=10, font=("Courier", 10), wrap="word")
        self.pred_details.pack(fill="both", expand=True, padx=20, pady=10)

    # ----- Results Tab -----
    def create_results_tab(self):
        tab = ttk.Frame(self.notebook)
        self.notebook.add(tab, text="📊 Results")

        tk.Label(tab, text="Model Performance Comparison", font=("Arial", 14, "bold")).pack(pady=10)
        columns = ("Model", "Preproc", "Accuracy", "Precision", "Recall", "F1", "BalAcc", "Kappa", "MCC", "AUC")
        self.results_tree = ttk.Treeview(tab, columns=columns, show="headings", height=8)
        for col in columns:
            self.results_tree.heading(col, text=col)
            self.results_tree.column(col, width=100)
        self.results_tree.pack(fill="x", padx=20, pady=10)

        cm_frame = tk.LabelFrame(tab, text="Confusion Matrix (Best Model)", font=("Arial", 12, "bold"))
        cm_frame.pack(fill="both", expand=True, padx=20, pady=10)
        self.cm_fig = Figure(figsize=(6, 4), dpi=100)
        self.cm_ax = self.cm_fig.add_subplot(111)
        self.cm_canvas = FigureCanvasTkAgg(self.cm_fig, master=cm_frame)
        self.cm_canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=10)

        roc_frame = tk.LabelFrame(tab, text="ROC Curves (Best Model)", font=("Arial", 12, "bold"))
        roc_frame.pack(fill="both", expand=True, padx=20, pady=10)
        self.roc_fig = Figure(figsize=(6, 4), dpi=100)
        self.roc_ax = self.roc_fig.add_subplot(111)
        self.roc_canvas = FigureCanvasTkAgg(self.roc_fig, master=roc_frame)
        self.roc_canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=10)

        report_frame = tk.LabelFrame(tab, text="Classification Report", font=("Arial", 12, "bold"))
        report_frame.pack(fill="both", expand=True, padx=20, pady=10)
        self.class_report_text = scrolledtext.ScrolledText(report_frame, height=8, font=("Courier", 9), wrap="word")
        self.class_report_text.pack(fill="both", expand=True, padx=10, pady=10)

    # ----- Pipeline Functions -----
    def train_models(self):
        if self.X_train is None:
            messagebox.showerror("Error", "Dataset not loaded!")
            return

        selected_models = [name for name, var in self.model_vars.items() if var.get()]
        if not selected_models:
            messagebox.showwarning("Warning", "Select at least one model!")
            return

        prep = self.prep_var.get()
        use_fs = self.feat_sel_var.get()
        self.progress.start()
        self.train_btn.config(state="disabled")
        self.status_var.set("Training...")
        self.log(f"\n{'=' * 60}")
        self.log(f"Preprocessing: {prep}")
        if use_fs:
            self.log("Feature Selection: Enabled (Top 15 features)")

        # Prepare data based on preprocessing
        if prep == "SMOTE":
            X_tr, y_tr = SMOTE(random_state=42).fit_resample(self.X_train, self.y_train)
        elif prep == "SMOTE+Tomek":
            X_tr, y_tr = SMOTETomek(random_state=42).fit_resample(self.X_train, self.y_train)
        else:
            X_tr, y_tr = self.X_train, self.y_train
        self.log(f"Training samples after resampling: {len(X_tr)}")

        scaler = StandardScaler()
        X_tr_scaled = scaler.fit_transform(X_tr)
        X_te_scaled = scaler.transform(self.X_test)

        # Feature Selection
        selector = None
        if use_fs:
            selector = SelectKBest(score_func=f_classif, k=15)
            X_tr_scaled = selector.fit_transform(X_tr_scaled, y_tr)
            X_te_scaled = selector.transform(X_te_scaled)
            self.feature_names = [self.X.columns[i] for i in selector.get_support(indices=True)]
        else:
            self.feature_names = self.X.columns.tolist()

        # Model definitions (Enhanced with Class Weights & Tuning)
        base_estimators = [
            ('rf', RandomForestClassifier(n_estimators=300, class_weight='balanced', random_state=42, n_jobs=-1)),
            ('et', ExtraTreesClassifier(n_estimators=300, class_weight='balanced', random_state=42, n_jobs=-1)),
            ('xgb', XGBClassifier(n_estimators=300, learning_rate=0.05, max_depth=5,
                                  subsample=0.8, colsample_bytree=0.8,
                                  objective="multi:softprob", num_class=3,
                                  random_state=42, eval_metric="mlogloss"))
        ]

        model_defs = {
            "Random Forest": base_estimators[0][1],
            "Extra Trees": base_estimators[1][1],
            "AdaBoost": AdaBoostClassifier(n_estimators=300, learning_rate=0.1, random_state=42),
            "Gradient Boosting": GradientBoostingClassifier(n_estimators=200, learning_rate=0.05, random_state=42),
            "XGBoost": base_estimators[2][1],
            "Stacking Ensemble": StackingClassifier(estimators=base_estimators,
                                                    final_estimator=LogisticRegression(max_iter=1000), cv=5, n_jobs=-1)
        }

        self.results = []
        self.models = {}

        def worker():
            try:
                for model_name in selected_models:
                    self.log(f"\nTraining {model_name}...")
                    model = model_defs[model_name]
                    start = time.time()

                    model.fit(X_tr_scaled, y_tr)

                    train_time = time.time() - start
                    y_pred = model.predict(X_te_scaled)

                    # Metrics
                    acc = accuracy_score(self.y_test, y_pred)
                    prec = precision_score(self.y_test, y_pred, average='weighted')
                    rec = recall_score(self.y_test, y_pred, average='weighted')
                    f1 = f1_score(self.y_test, y_pred, average='weighted')
                    bal_acc = balanced_accuracy_score(self.y_test, y_pred)
                    kappa = cohen_kappa_score(self.y_test, y_pred)
                    mcc = matthews_corrcoef(self.y_test, y_pred)

                    # AUC
                    if hasattr(model, "predict_proba"):
                        y_score = model.predict_proba(X_te_scaled)
                        y_test_bin = label_binarize(self.y_test, classes=[0, 1, 2])
                        auc = roc_auc_score(y_test_bin, y_score, multi_class="ovr")
                    else:
                        auc = np.nan

                    result = {
                        "Model": model_name,
                        "Preprocessing": prep,
                        "Accuracy": acc,
                        "Precision": prec,
                        "Recall": rec,
                        "F1": f1,
                        "Balanced_Accuracy": bal_acc,
                        "Kappa": kappa,
                        "MCC": mcc,
                        "AUC": auc,
                        "Train_Time": train_time,
                        "y_pred": y_pred,
                        "model": model,
                        "scaler": scaler,
                        "selector": selector
                    }
                    self.results.append(result)
                    self.models[model_name] = model

                    self.log(f"  Test Accuracy: {acc:.4f}  F1: {f1:.4f}  AUC: {auc:.4f}")

                # Summary: best model
                if self.results:
                    best = max(self.results, key=lambda x: x["Accuracy"])
                    self.log(f"\n{'=' * 60}")
                    self.log("BEST MODEL:")
                    self.log(f"  {best['Model']} with {best['Preprocessing']}")
                    self.log(f"  Test Accuracy: {best['Accuracy']:.4f} ({best['Accuracy'] * 100:.2f}%)")
                    self.log(f"  F1-Score: {best['F1']:.4f}")
                    self.log(f"  MCC: {best['MCC']:.4f}")
                    self.log(f"{'=' * 60}")

                    self.log("\nClassification Report (best model):")
                    report = classification_report(self.y_test, best["y_pred"],
                                                   target_names=['Normal', 'Suspect', 'Pathologic'],
                                                   zero_division=0)
                    self.log(report)

                # Update UI
                self.root.after(0, self.update_results)
                self.root.after(0, self.update_dashboard)
                self.root.after(0, self.update_feature_importance)
                self.root.after(0, self.update_results_tree)

                # Update Predict tab models dynamically
                self.root.after(0, lambda: self.pred_combo.config(values=list(self.models.keys())))

                self.status_var.set("Training complete")
                self.log("✓ Training finished.")
            except Exception as e:
                self.root.after(0, lambda: messagebox.showerror("Error", str(e)))
            finally:
                self.root.after(0, self.stop_training)

        Thread(target=worker).start()

    def stop_training(self):
        self.progress.stop()
        self.train_btn.config(state="normal")

    # ----- UI Updates -----
    def update_results(self):
        if not self.results:
            return
        best = max(self.results, key=lambda x: x["Accuracy"])
        self.cm_ax.clear()
        cm = confusion_matrix(self.y_test, best["y_pred"])
        sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", ax=self.cm_ax,
                    xticklabels=self.class_names, yticklabels=self.class_names)
        self.cm_ax.set_title(f"Confusion Matrix – {best['Model']}\nAcc: {best['Accuracy']:.2%}")
        self.cm_ax.set_xlabel("Predicted")
        self.cm_ax.set_ylabel("Actual")
        self.cm_canvas.draw()

        self.roc_ax.clear()
        if hasattr(best["model"], "predict_proba"):
            # Reapply transformations for ROC curve calculation
            X_test_scaled = best["scaler"].transform(self.X_test)
            if best["selector"] is not None:
                X_test_scaled = best["selector"].transform(X_test_scaled)

            y_proba = best["model"].predict_proba(X_test_scaled)
            for i, cls in enumerate(self.class_names):
                fpr, tpr, _ = roc_curve((self.y_test == i).astype(int), y_proba[:, i])
                auc_score = roc_auc_score((self.y_test == i).astype(int), y_proba[:, i])
                self.roc_ax.plot(fpr, tpr, label=f"{cls} (AUC={auc_score:.3f})")
            self.roc_ax.plot([0, 1], [0, 1], "k--", alpha=0.3, label="Random")
            self.roc_ax.set_xlabel("False Positive Rate")
            self.roc_ax.set_ylabel("True Positive Rate")
            self.roc_ax.set_title("ROC Curves (One-vs-Rest)")
            self.roc_ax.legend(loc="lower right")
            self.roc_ax.grid(True, alpha=0.3)
        self.roc_canvas.draw()

        report = classification_report(self.y_test, best["y_pred"],
                                       target_names=self.class_names, zero_division=0)
        self.class_report_text.delete(1.0, tk.END)
        self.class_report_text.insert(tk.END, report)

    def update_results_tree(self):
        for item in self.results_tree.get_children():
            self.results_tree.delete(item)
        # Sort results by Accuracy descending
        sorted_results = sorted(self.results, key=lambda x: x["Accuracy"], reverse=True)
        for res in sorted_results:
            self.results_tree.insert("", "end", values=(
                res["Model"],
                res["Preprocessing"],
                f"{res['Accuracy']:.4f}",
                f"{res['Precision']:.4f}",
                f"{res['Recall']:.4f}",
                f"{res['F1']:.4f}",
                f"{res['Balanced_Accuracy']:.4f}",
                f"{res['Kappa']:.4f}",
                f"{res['MCC']:.4f}",
                f"{res['AUC']:.4f}"
            ))

    def update_dashboard(self):
        if self.df is None:
            return
        self.metric_cards["Total Samples"].config(text=f"{len(self.df):,}")
        self.metric_cards["Features"].config(text=f"{self.X.shape[1]}")
        nsp = self.df["NSP"].value_counts()
        self.metric_cards["Normal"].config(text=f"{nsp.get(1, 0):,}")
        self.metric_cards["Suspect"].config(text=f"{nsp.get(2, 0):,}")
        self.metric_cards["Pathologic"].config(text=f"{nsp.get(3, 0):,}")

        if self.results:
            best = max(self.results, key=lambda x: x["Accuracy"])
            self.metric_cards["Best Accuracy"].config(text=f"{best['Accuracy'] * 100:.1f}%")
            self.metric_cards["Best F1"].config(text=f"{best['F1']:.4f}")
            self.metric_cards["Best MCC"].config(text=f"{best['MCC']:.4f}")

        self.dash_ax.clear()
        if self.results:
            models = [r["Model"] for r in self.results]
            accs = [r["Accuracy"] for r in self.results]
            colors = plt.cm.Set3(np.linspace(0, 1, len(models)))
            self.dash_ax.bar(models, accs, color=colors)
            self.dash_ax.set_ylim([0, 1])
            self.dash_ax.set_ylabel("Accuracy")
            self.dash_ax.set_title("Model Accuracy Comparison")
            self.dash_ax.tick_params(axis='x', rotation=45)
        self.dash_canvas.draw()

    def update_feature_importance(self):
        self.fi_ax.clear()
        if "Random Forest" in self.models:
            rf = self.models["Random Forest"]
            if hasattr(rf, "feature_importances_"):
                imp = rf.feature_importances_
                # Match importance array with used feature names
                feat_df = pd.DataFrame({"feature": self.feature_names, "importance": imp})
                feat_df = feat_df.sort_values("importance", ascending=True).tail(15)
                colors = plt.cm.YlOrRd(np.linspace(0.3, 0.9, len(feat_df)))
                self.fi_ax.barh(feat_df["feature"], feat_df["importance"], color=colors)
                self.fi_ax.set_xlabel("Importance")
                self.fi_ax.set_title("Top Features (Random Forest)")
                self.fi_ax.grid(True, alpha=0.3, axis="x")
        else:
            self.fi_ax.text(0.5, 0.5, "Train Random Forest to see importance", ha="center", va="center")
        self.fi_canvas.draw()

    # ----- Prediction -----
    def populate_feature_entries(self):
        for widget in self.feature_frame.winfo_children():
            widget.destroy()
        self.feature_entries = {}
        if self.X is None:
            return
        for feature in self.X.columns:
            frm = tk.Frame(self.feature_frame, bg="white")
            frm.pack(fill="x", pady=2)
            tk.Label(frm, text=f"{feature}:", width=22, anchor="w", bg="white",
                     font=("Arial", 9)).pack(side="left")
            entry = tk.Entry(frm, width=15)
            entry.pack(side="left", padx=5)
            entry.insert(0, "0")
            self.feature_entries[feature] = entry

    def load_sample(self, event=None):
        sample_type = self.sample_var.get()
        class_map = {"Normal": 0, "Suspect": 1, "Pathologic": 2}
        if self.X is None:
            return
        if sample_type in class_map:
            samples = self.X[self.y == class_map[sample_type]]
            sample = samples.iloc[0] if len(samples) > 0 else self.X.iloc[0]
        elif sample_type == "Random":
            sample = self.X.iloc[np.random.randint(0, len(self.X))]
        else:
            return
        for feature in self.X.columns:
            if feature in self.feature_entries:
                self.feature_entries[feature].delete(0, tk.END)
                self.feature_entries[feature].insert(0, f"{sample[feature]:.3f}")

    def make_prediction(self):
        if not self.models:
            messagebox.showwarning("Warning", "Train a model first!")
            return
        model_name = self.pred_model_var.get()
        if model_name not in self.models:
            messagebox.showerror("Error", f"Model {model_name} not trained!")
            return

        model = self.models[model_name]
        scaler = None
        selector = None

        for res in self.results:
            if res["Model"] == model_name:
                scaler = res["scaler"]
                selector = res["selector"]
                break

        if scaler is None:
            messagebox.showerror("Error", "Scaler not found for this model.")
            return

        try:
            features = [float(self.feature_entries[f].get()) for f in self.X.columns]
            features_scaled = scaler.transform([features])

            if selector is not None:
                features_scaled = selector.transform(features_scaled)

            pred = model.predict(features_scaled)[0]
            proba = model.predict_proba(features_scaled)[0] if hasattr(model, "predict_proba") else [0, 0, 0]
            cls_name = self.class_names[pred]
            conf = proba[pred] * 100 if len(proba) > 0 else 100
            colors = {0: "#2ecc71", 1: "#f39c12", 2: "#e74c3c"}
            self.pred_label.config(text=f"{cls_name} – {conf:.1f}% confidence", fg=colors[pred])
            self.pred_result_frame.config(bg="#f8f9fa")
            self.pred_details.delete(1.0, tk.END)
            self.pred_details.insert(tk.END, f"Model: {model_name}\nPredicted: {cls_name}\nConfidence: {conf:.1f}%\n\n")
            self.pred_details.insert(tk.END, "Class Probabilities:\n")
            for i, name in enumerate(self.class_names):
                self.pred_details.insert(tk.END, f"  {name}: {proba[i] * 100:.1f}%\n")
        except Exception as e:
            messagebox.showerror("Error", str(e))

    # ----- Report Generation -----
    def generate_word_report(self):
        if not DOCX_AVAILABLE:
            messagebox.showerror("Error", "python-docx not installed.")
            return
        if not self.results:
            messagebox.showwarning("Warning", "No results to report.")
            return

        try:
            doc = Document()
            title = doc.add_heading("CTG ML Pipeline Report", 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

            doc.add_heading("1. Dataset Overview", level=1)
            table = doc.add_table(rows=5, cols=2)
            table.style = "Table Grid"
            nsp = self.df["NSP"].value_counts()
            rows_data = [
                ("Total Samples", str(len(self.df))),
                ("Features", str(self.X.shape[1])),
                ("Target", "NSP (Normal, Suspect, Pathologic)"),
                ("Class Distribution",
                 f"Normal: {nsp.get(1, 0)}, Suspect: {nsp.get(2, 0)}, Pathologic: {nsp.get(3, 0)}"),
                ("Train/Test Split", f"{len(self.X_train)} / {len(self.X_test)}")
            ]
            for i, (k, v) in enumerate(rows_data):
                table.rows[i].cells[0].text = k
                table.rows[i].cells[1].text = v

            doc.add_heading("2. Model Performance", level=1)
            table = doc.add_table(rows=len(self.results) + 1, cols=10)
            table.style = "Table Grid"
            headers = ["Model", "Preproc", "Acc", "Prec", "Recall", "F1", "BalAcc", "Kappa", "MCC", "AUC"]
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h

            sorted_results = sorted(self.results, key=lambda x: x["Accuracy"], reverse=True)
            for i, r in enumerate(sorted_results):
                row = table.rows[i + 1]
                row.cells[0].text = r["Model"]
                row.cells[1].text = r["Preprocessing"]
                row.cells[2].text = f"{r['Accuracy']:.4f}"
                row.cells[3].text = f"{r['Precision']:.4f}"
                row.cells[4].text = f"{r['Recall']:.4f}"
                row.cells[5].text = f"{r['F1']:.4f}"
                row.cells[6].text = f"{r['Balanced_Accuracy']:.4f}"
                row.cells[7].text = f"{r['Kappa']:.4f}"
                row.cells[8].text = f"{r['MCC']:.4f}"
                row.cells[9].text = f"{r['AUC']:.4f}"

            best = sorted_results[0]
            doc.add_heading("3. Best Model Details", level=1)
            doc.add_paragraph(f"Model: {best['Model']}")
            doc.add_paragraph(f"Preprocessing: {best['Preprocessing']}")
            doc.add_paragraph(f"Accuracy: {best['Accuracy']:.2%}")
            doc.add_paragraph(f"F1-Score: {best['F1']:.4f}")
            doc.add_paragraph(f"MCC: {best['MCC']:.4f}")

            doc.add_heading("4. Confusion Matrix", level=1)
            fig, ax = plt.subplots(figsize=(6, 4))
            cm = confusion_matrix(self.y_test, best["y_pred"])
            sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", ax=ax,
                        xticklabels=self.class_names, yticklabels=self.class_names)
            ax.set_title(f"Confusion Matrix ({best['Model']})")
            img_buffer = BytesIO()
            fig.savefig(img_buffer, format="png", dpi=150, bbox_inches="tight")
            img_buffer.seek(0)
            doc.add_picture(img_buffer, width=Inches(5))
            plt.close(fig)

            doc.add_heading("5. Classification Report", level=1)
            report = classification_report(self.y_test, best["y_pred"],
                                           target_names=self.class_names, zero_division=0)
            p = doc.add_paragraph()
            run = p.add_run(report)
            run.font.name = "Courier New"
            run.font.size = Pt(9)

            filename = filedialog.asksaveasfilename(
                title="Save Report", defaultextension=".docx",
                filetypes=[("Word documents", "*.docx")]
            )
            if filename:
                doc.save(filename)
                messagebox.showinfo("Success", f"Report saved to {filename}")

        except Exception as e:
            messagebox.showerror("Error", str(e))

    # ----- Helpers -----
    def log(self, msg):
        self.log_text.insert(tk.END, msg + "\n")
        self.log_text.see(tk.END)
        self.root.update()

    def load_data_dialog(self):
        path = filedialog.askopenfilename(filetypes=[("CSV files", "*.csv")])
        if path:
            self.load_data(path)


if __name__ == "__main__":
    root = tk.Tk()
    app = CTGPipelineGUI(root)
    root.mainloop()